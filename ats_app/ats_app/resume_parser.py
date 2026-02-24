"""
Resume Parser Module - File Upload and Text Extraction
Supports PDF, DOCX, and TXT file formats
"""
import io
import os
import logging
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Tuple, Optional, List

from config import settings

logger = logging.getLogger(__name__)

# PDF parsing
try:
    from pypdf import PdfReader
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

# DOCX parsing
try:
    from docx import Document
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

# Upload directory configuration
UPLOAD_DIR = Path(settings.UPLOAD_DIR)


def ensure_upload_dir() -> Path:
    """Create upload directory if it doesn't exist"""
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    return UPLOAD_DIR


def get_supported_extensions() -> List[str]:
    """Return list of supported file extensions"""
    extensions = [".txt"]
    if PDF_AVAILABLE:
        extensions.append(".pdf")
    if DOCX_AVAILABLE:
        extensions.append(".docx")
    return extensions


def validate_file_magic(file_bytes: bytes, expected_ext: str) -> Tuple[bool, Optional[str]]:
    """
    Validate file magic bytes match the declared extension.

    Args:
        file_bytes: Raw file bytes to validate
        expected_ext: Expected file extension (e.g., '.pdf', '.docx', '.txt')

    Returns:
        Tuple of (is_valid, error_message)
        If valid, error_message is None
    """
    if not file_bytes:
        return False, "No file data provided"

    if len(file_bytes) < 4:
        return False, "File too small to validate"

    ext = expected_ext.lower()

    # PDF validation: should start with %PDF
    if ext == '.pdf':
        if file_bytes[:4] == b'%PDF':
            return True, None
        return False, "File does not appear to be a valid PDF (magic bytes mismatch)"

    # DOCX validation: should start with PK\x03\x04 (ZIP format)
    elif ext == '.docx':
        if file_bytes[:4] == b'PK\x03\x04':
            return True, None
        return False, "File does not appear to be a valid DOCX (magic bytes mismatch)"

    # TXT validation: should NOT match PDF, DOCX, or other binary formats
    elif ext == '.txt':
        # Check if it starts with known binary signatures
        if file_bytes[:4] == b'%PDF':
            return False, "File appears to be a PDF, not a text file"
        if file_bytes[:4] == b'PK\x03\x04':
            return False, "File appears to be a ZIP/DOCX, not a text file"
        if file_bytes[:2] == b'MZ':
            return False, "File appears to be an executable, not a text file"
        if file_bytes[:4] == b'\x7fELF':
            return False, "File appears to be an ELF binary, not a text file"
        if file_bytes[:2] == b'\xff\xd8':
            return False, "File appears to be a JPEG image, not a text file"
        if file_bytes[:8] == b'\x89PNG\r\n\x1a\n':
            return False, "File appears to be a PNG image, not a text file"
        # Text files can have various content, so we accept it if it's not binary
        return True, None

    else:
        return False, f"Unsupported file extension: {ext}"


def check_zip_bomb(file_bytes: bytes) -> Tuple[bool, Optional[str]]:
    """
    Check if a ZIP file (DOCX) might be a zip bomb.

    Args:
        file_bytes: Raw file bytes of ZIP/DOCX file

    Returns:
        Tuple of (is_safe, error_message)
        If safe, error_message is None
    """
    max_decompressed_mb = settings.MAX_DECOMPRESSED_SIZE_MB
    max_decompressed_bytes = max_decompressed_mb * 1024 * 1024

    try:
        zip_stream = io.BytesIO(file_bytes)
        with zipfile.ZipFile(zip_stream, 'r') as zip_ref:
            total_decompressed = 0

            for info in zip_ref.infolist():
                total_decompressed += info.file_size

                if total_decompressed > max_decompressed_bytes:
                    return False, f"Decompressed size exceeds limit ({max_decompressed_mb}MB) - possible zip bomb"

            return True, None

    except zipfile.BadZipFile:
        return False, "Invalid ZIP file structure"
    except Exception as e:
        logger.warning(f"Error checking zip bomb: {e}")
        return False, f"Error validating ZIP file: {str(e)}"


def scan_file_content(text: str) -> Tuple[bool, List[str]]:
    """
    Scan extracted text for suspicious patterns.

    Args:
        text: Extracted text content

    Returns:
        Tuple of (is_clean, warnings)
        If suspicious content is found, is_clean is False
    """
    warnings = []

    if not text:
        return True, warnings

    # Check for extremely long lines (might indicate binary data)
    lines = text.split('\n')
    max_line_length = 0
    long_line_count = 0

    for line in lines:
        line_len = len(line)
        if line_len > max_line_length:
            max_line_length = line_len
        if line_len > 10000:
            long_line_count += 1

    if long_line_count > 0:
        warnings.append(f"Found {long_line_count} extremely long line(s) (>10000 chars) - possible binary data")

    # Check ratio of non-printable characters
    if len(text) > 0:
        printable_count = sum(1 for c in text if c.isprintable() or c in '\n\r\t')
        printable_ratio = printable_count / len(text)

        if printable_ratio < 0.8:
            warnings.append(f"High ratio of non-printable characters ({100-printable_ratio*100:.1f}%) - possible binary data")

    # Check for null bytes (should not be in text)
    if '\x00' in text:
        warnings.append("Null bytes detected in extracted text - possible binary data")

    # If we have warnings, consider it suspicious but not necessarily unsafe
    # (some PDFs have metadata that triggers these)
    is_clean = len(warnings) == 0

    return is_clean, warnings


def parse_pdf(file_bytes: bytes) -> Tuple[str, List[str]]:
    """
    Extract text from PDF file bytes.
    Returns (extracted_text, warnings_list)
    """
    if not PDF_AVAILABLE:
        return "", ["PDF parsing not available - pypdf not installed"]

    warnings = []
    text_parts = []

    try:
        pdf_stream = io.BytesIO(file_bytes)
        reader = PdfReader(pdf_stream)

        if len(reader.pages) == 0:
            return "", ["PDF file appears to be empty"]

        # Enforce max pages limit for security
        max_pages = settings.MAX_PDF_PAGES
        page_count = len(reader.pages)

        if page_count > max_pages:
            return "", [f"PDF exceeds maximum page limit ({page_count} pages, max {max_pages} allowed)"]

        for page_num, page in enumerate(reader.pages, 1):
            try:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
                else:
                    warnings.append(f"Page {page_num}: No text extracted (may be scanned/image)")
            except (ValueError, AttributeError) as e:
                logger.warning(f"Failed to extract text from PDF page {page_num}: {e}")
                warnings.append(f"Page {page_num}: Error extracting text - {str(e)}")

        if not text_parts:
            warnings.append("No text could be extracted from PDF - may be scanned or image-based")
            return "", warnings

        extracted_text = "\n\n".join(text_parts)

        # Clean up common PDF extraction artifacts
        extracted_text = clean_extracted_text(extracted_text)

        return extracted_text, warnings

    except (IOError, OSError) as e:
        logger.error(f"I/O error reading PDF file: {e}")
        return "", [f"Failed to parse PDF: {str(e)}"]
    except Exception as e:
        logger.exception("Unexpected error parsing PDF file")
        return "", [f"Failed to parse PDF: {str(e)}"]


def parse_docx(file_bytes: bytes) -> Tuple[str, List[str]]:
    """
    Extract text from DOCX file bytes.
    Returns (extracted_text, warnings_list)
    """
    if not DOCX_AVAILABLE:
        return "", ["DOCX parsing not available - python-docx not installed"]

    # Check for zip bomb before parsing
    is_safe, error_msg = check_zip_bomb(file_bytes)
    if not is_safe:
        return "", [error_msg]

    warnings = []
    text_parts = []

    try:
        docx_stream = io.BytesIO(file_bytes)
        doc = Document(docx_stream)

        # Extract paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)

        # Extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    if cell.text.strip():
                        row_text.append(cell.text.strip())
                if row_text:
                    text_parts.append(" | ".join(row_text))

        if not text_parts:
            warnings.append("Document appears to be empty or contains no extractable text")
            return "", warnings

        extracted_text = "\n".join(text_parts)
        extracted_text = clean_extracted_text(extracted_text)

        return extracted_text, warnings

    except (IOError, OSError) as e:
        logger.error(f"I/O error reading DOCX file: {e}")
        return "", [f"Failed to parse DOCX: {str(e)}"]
    except Exception as e:
        logger.exception("Unexpected error parsing DOCX file")
        return "", [f"Failed to parse DOCX: {str(e)}"]


def parse_txt(file_bytes: bytes) -> Tuple[str, List[str]]:
    """
    Extract text from TXT file bytes.
    Returns (extracted_text, warnings_list)
    """
    warnings = []

    # Try different encodings
    encodings = ['utf-8', 'latin-1', 'cp1252', 'ascii']

    for encoding in encodings:
        try:
            text = file_bytes.decode(encoding)
            if text.strip():
                return clean_extracted_text(text), warnings
        except (UnicodeDecodeError, LookupError):
            continue

    return "", ["Could not decode text file - unsupported encoding"]


def clean_extracted_text(text: str) -> str:
    """Clean up extracted text by removing common artifacts"""
    if not text:
        return text

    # Replace multiple spaces with single space
    import re
    text = re.sub(r' +', ' ', text)

    # Replace multiple newlines with double newline
    text = re.sub(r'\n{3,}', '\n\n', text)

    # Remove leading/trailing whitespace from each line
    lines = [line.strip() for line in text.split('\n')]
    text = '\n'.join(lines)

    # Remove leading/trailing whitespace from whole text
    text = text.strip()

    return text


def parse_resume(file_bytes: bytes, filename: str) -> Tuple[str, List[str]]:
    """
    Parse resume file and extract text.

    Args:
        file_bytes: Raw file bytes
        filename: Original filename (used to determine file type)

    Returns:
        Tuple of (extracted_text, warnings_list)
        If extraction fails, text will be empty and warnings will explain why
    """
    if not filename:
        return "", ["No filename provided"]

    # Get file extension
    ext = Path(filename).suffix.lower()

    # Validate file type is supported
    if ext not in get_supported_extensions():
        supported = ", ".join(get_supported_extensions())
        return "", [f"Unsupported file type: {ext}. Supported formats: {supported}"]

    # Validate magic bytes match declared extension
    is_valid_magic, magic_error = validate_file_magic(file_bytes, ext)
    if not is_valid_magic:
        return "", [magic_error]

    # Parse based on file type
    if ext == '.pdf':
        extracted_text, parse_warnings = parse_pdf(file_bytes)
    elif ext == '.docx':
        extracted_text, parse_warnings = parse_docx(file_bytes)
    elif ext == '.txt':
        extracted_text, parse_warnings = parse_txt(file_bytes)
    else:
        supported = ", ".join(get_supported_extensions())
        return "", [f"Unsupported file type: {ext}. Supported formats: {supported}"]

    # If parsing failed, return early
    if not extracted_text:
        return extracted_text, parse_warnings

    # Scan extracted content for suspicious patterns
    is_clean, content_warnings = scan_file_content(extracted_text)

    # Combine warnings
    all_warnings = parse_warnings + content_warnings

    return extracted_text, all_warnings


def save_resume_file(candidate_id: int, file_bytes: bytes, filename: str) -> Tuple[Optional[str], Optional[str]]:
    """
    Save uploaded resume file to disk.

    Args:
        candidate_id: ID of the candidate
        file_bytes: Raw file bytes
        filename: Original filename

    Returns:
        Tuple of (saved_file_path, error_message)
        If save fails, path will be None and error will explain why
    """
    if not file_bytes:
        return None, "No file data provided"

    if not filename:
        return None, "No filename provided"

    # Check for null bytes in filename
    if '\x00' in filename:
        return None, "Invalid filename: contains null bytes"

    # Check filename length
    if len(filename) > settings.MAX_FILENAME_LENGTH:
        return None, f"Filename too long (max {settings.MAX_FILENAME_LENGTH} characters)"

    try:
        ensure_upload_dir()

        # Generate unique filename: {candidate_id}_{timestamp}_{original_filename}
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

        # Clean filename to prevent path traversal
        safe_filename = Path(filename).name  # Remove any directory components
        safe_filename = "".join(c for c in safe_filename if c.isalnum() or c in '.-_')

        if not safe_filename:
            safe_filename = "resume"

        # Preserve extension
        original_ext = Path(filename).suffix.lower()
        if original_ext and not safe_filename.endswith(original_ext):
            safe_filename += original_ext

        unique_filename = f"{candidate_id}_{timestamp}_{safe_filename}"

        # Enforce filename length limit on final filename too
        if len(unique_filename) > settings.MAX_FILENAME_LENGTH:
            # Truncate safe_filename portion to fit
            max_safe_len = settings.MAX_FILENAME_LENGTH - len(f"{candidate_id}_{timestamp}_") - len(original_ext)
            if max_safe_len > 0:
                safe_filename = safe_filename[:max_safe_len] + original_ext
                unique_filename = f"{candidate_id}_{timestamp}_{safe_filename}"
            else:
                return None, "Filename too long even after sanitization"

        file_path = UPLOAD_DIR / unique_filename

        # Verify resolved path is still within UPLOAD_DIR (prevents path traversal)
        try:
            resolved_path = file_path.resolve()
            resolved_upload_dir = UPLOAD_DIR.resolve()

            # Check if the resolved path is within the upload directory
            if not str(resolved_path).startswith(str(resolved_upload_dir)):
                logger.error(f"Path traversal attempt detected: {filename}")
                return None, "Invalid file path"

        except (ValueError, OSError) as e:
            logger.error(f"Path resolution failed for {filename}: {e}")
            return None, "Invalid file path"

        # Write file
        with open(file_path, 'wb') as f:
            f.write(file_bytes)

        return str(file_path), None

    except (IOError, OSError) as e:
        logger.error(f"Failed to save resume file for candidate {candidate_id}: {e}")
        return None, f"Failed to save file: {str(e)}"
    except Exception as e:
        logger.exception(f"Unexpected error saving resume file for candidate {candidate_id}")
        return None, f"Failed to save file: {str(e)}"


def get_resume_file(resume_path: str) -> Tuple[Optional[bytes], Optional[str]]:
    """
    Retrieve resume file from disk.

    Args:
        resume_path: Path to the saved resume file

    Returns:
        Tuple of (file_bytes, error_message)
        If retrieval fails, bytes will be None and error will explain why
    """
    if not resume_path:
        return None, "No file path provided"

    try:
        file_path = Path(resume_path)

        if not file_path.exists():
            return None, "File not found"

        if not file_path.is_file():
            return None, "Path is not a file"

        with open(file_path, 'rb') as f:
            return f.read(), None

    except (IOError, OSError) as e:
        logger.error(f"Failed to read resume file at {resume_path}: {e}")
        return None, f"Failed to read file: {str(e)}"
    except Exception as e:
        logger.exception(f"Unexpected error reading resume file at {resume_path}")
        return None, f"Failed to read file: {str(e)}"


def delete_resume_file(resume_path: str) -> Optional[str]:
    """
    Delete resume file from disk.

    Args:
        resume_path: Path to the saved resume file

    Returns:
        Error message if deletion fails, None if successful
    """
    if not resume_path:
        return None  # No file to delete

    try:
        file_path = Path(resume_path)

        if file_path.exists():
            file_path.unlink()

        return None

    except (IOError, OSError) as e:
        logger.error(f"Failed to delete resume file at {resume_path}: {e}")
        return f"Failed to delete file: {str(e)}"
    except Exception as e:
        logger.exception(f"Unexpected error deleting resume file at {resume_path}")
        return f"Failed to delete file: {str(e)}"


def validate_file_size(file_bytes, max_size_mb: float = None) -> Optional[str]:
    """
    Validate file size is within limits.

    Args:
        file_bytes: Raw file bytes or a Streamlit UploadedFile object
        max_size_mb: Maximum allowed size in megabytes (default from settings)

    Returns:
        Error message if file is too large, None if OK
    """
    if not file_bytes:
        return "No file data"

    if max_size_mb is None:
        max_size_mb = settings.MAX_RESUME_SIZE_MB

    # Support both raw bytes and Streamlit UploadedFile objects
    if hasattr(file_bytes, 'size'):
        size_mb = file_bytes.size / (1024 * 1024)
    else:
        size_mb = len(file_bytes) / (1024 * 1024)

    if size_mb > max_size_mb:
        return f"File too large: {size_mb:.1f}MB (max {max_size_mb:.1f}MB)"

    return None


def get_parse_status() -> dict:
    """Return status of parsing libraries"""
    return {
        "pdf": PDF_AVAILABLE,
        "docx": DOCX_AVAILABLE,
        "txt": True,
        "supported_extensions": get_supported_extensions()
    }
